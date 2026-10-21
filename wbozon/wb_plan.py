"""Отчёт WB для плана: текущие продажи и остатки одного артикула + прошлогодняя динамика других.

    python run.py --wb-plan                      # БЖ5 сейчас; БЖ2, БЖ4, БЖ11 за 20.10–31.12.2025

Все даты — по московскому времени (статистика WB и отчёты отдают время по Москве).
Источники:
* заказы и отмены по размерам — statistics-api /api/v1/supplier/orders (WB хранит их ~6 месяцев);
* продажи и возвраты по размерам — /api/v1/supplier/sales, а раньше — детализация отчёта о реализации
  (finance-api /api/finance/v1/sales-reports/detailed, дата продажи saleDt);
* заказы, выкупы и отмены по дням по артикулу целиком (без размеров) — отчёт «Воронка продаж»
  DETAIL_HISTORY_REPORT (seller-analytics-api /api/v2/nm-report/downloads), не дальше 1 года назад;
* история остатков по дням и размерам — STOCK_HISTORY_DAILY_CSV, не дальше 3 месяцев назад;
* остатки сейчас — /api/analytics/v1/stocks-report/wb-warehouses и /seller-warehouses.
"""
import datetime as dt
import io
import json
import re
import time
import uuid
import zipfile

import pandas as pd
from openpyxl.utils import get_column_letter

from . import wb
from .common import DATA, RAW, log, read_csv, request, yesterday

ANALYTICS = wb.ANALYTICS
CUR = "БЖ5"
PAST = ["БЖ2", "БЖ4", "БЖ11"]
PAST_FROM, PAST_TO = dt.date(2025, 10, 20), dt.date(2025, 12, 31)


def _size_key(s):
    m = re.match(r"\d+", str(s))
    return int(m.group()) if m else 0


def _sort(df, cols):
    return df.sort_values(cols, key=lambda s: s.map(_size_key) if s.name == "Размер" else s, kind="stable")


# ---------- отчёты seller-analytics CSV ----------

def nm_report(report_type, params):
    """Создаёт CSV-отчёт, ждёт готовности, возвращает DataFrame. Кэш в data/raw/wb_reports/."""
    rid = str(uuid.uuid4())
    request("POST", f"{ANALYTICS}/api/v2/nm-report/downloads", headers=wb._headers(),
            body={"id": rid, "reportType": report_type, "params": params}, pause=21)
    for _ in range(60):
        lst = request("GET", f"{ANALYTICS}/api/v2/nm-report/downloads", headers=wb._headers(),
                      params={"filter[downloadIds]": rid}) or {}
        st = next((x for x in lst.get("data") or [] if x["id"] == rid), {})
        if st.get("status") == "SUCCESS":
            break
        if st.get("status") == "FAILED":
            raise wb.ApiError(f"отчёт {report_type} не сформирован")
        time.sleep(10)
    import requests
    r = requests.get(f"{ANALYTICS}/api/v2/nm-report/downloads/file/{rid}", headers=wb._headers(), timeout=120)
    r.raise_for_status()
    z = zipfile.ZipFile(io.BytesIO(r.content))
    df = pd.read_csv(z.open(z.namelist()[0]), dtype=str, keep_default_na=False)
    (RAW / "wb_reports").mkdir(parents=True, exist_ok=True)
    df.to_csv(RAW / "wb_reports" / f"{report_type}_{dt.date.today()}_{rid[:8]}.csv", index=False, encoding="utf-8-sig")
    return df


def _nm_ids(prices, arts):
    return {g["vendorCode"].upper(): g["nmID"] for g in prices if g["vendorCode"].upper() in arts}


# ---------- БЖ5 ----------

def order_feed(nm, sidx):
    """Лента заказов (/api/analytics/v1/order-feed): все оформленные заказы за последние 31 день,
    включая те, которых нет в statistics/orders (неподтверждённая оплата). Возвращает
    (DataFrame[date, size, status, qty, srid], первая полностью покрытая дата)."""
    start = dt.datetime.now(dt.timezone.utc) - dt.timedelta(days=30, hours=23)
    rows, offset, snap = [], 0, None
    while True:
        body = {"selectedPeriod": {"start": start.strftime("%Y-%m-%dT%H:%M:%SZ")}, "nmIds": [nm],
                "pagination": {"limit": 1000, "offset": offset, **({"snapshotTime": snap} if snap else {})}}
        d = (request("POST", f"{ANALYTICS}/api/analytics/v1/order-feed", headers=wb._headers(), body=body, pause=61) or {}).get("data") or {}
        part = d.get("orders") or []
        rows += part
        snap = d.get("snapshotTime") or snap
        if len(part) < 1000:
            break
        offset += len(part)
    f = pd.DataFrame(rows)
    if not len(f):
        return f, None
    f["date"] = (pd.to_datetime(f["createdAt"]) + pd.Timedelta(hours=3)).dt.strftime("%Y-%m-%d")
    f["size"] = f["chrtId"].map(lambda c: sidx.get(int(c), (0, "", str(c)))[2])
    f["status"] = f["status"].map(lambda x: "cancelled" if x == "cancel" else "ordered")
    f["qty"] = 1
    first_full = ((start + dt.timedelta(hours=3)).date() + dt.timedelta(days=1)).isoformat()
    return f[["date", "size", "status", "qty", "srid"]], first_full


def orders_by_size(art, feed, feed_from):
    """Заказы по размерам: statistics/orders до feed_from, лента заказов начиная с feed_from."""
    o = read_csv(DATA / "wb_orders.csv")
    o = o[o["article"].str.upper() == art][["date", "size", "status", "qty", "srid"]].copy()
    o["qty"] = pd.to_numeric(o["qty"])
    if feed is not None and len(feed) and feed_from:
        o = pd.concat([o[o["date"] < feed_from], feed[feed["date"] >= feed_from]], ignore_index=True)
    return o[o["date"] <= yesterday().isoformat()]

def cur_orders_sheet(art, funnel, start, o, feed_from):
    s = read_csv(DATA / "wb_sales.csv")
    s = s[s["article"].str.upper() == art].copy()
    end = yesterday().isoformat()
    s["qty"] = pd.to_numeric(s["qty"])
    od = o[o["date"] <= end].pivot_table(index=["date", "size"], columns="status", values="qty", aggfunc="sum", fill_value=0)
    od["Заказано, шт"] = od.sum(axis=1)
    od = od.rename(columns={"cancelled": "Отменено, шт"}).drop(columns=["ordered"], errors="ignore")
    sd = s[(s["date"] >= start) & (s["date"] <= end)].pivot_table(index=["date", "size"], columns="status",
                                                                     values="qty", aggfunc="sum", fill_value=0)
    sd = sd.rename(columns={"sale": "Продано, шт", "return": "Возврат, шт"})
    df = od.join(sd, how="outer").reset_index().rename(columns={"date": "Дата", "size": "Размер"})
    df["Источник"] = df["Дата"].map(lambda d: "по размерам: заказы — " + ("лента заказов" if feed_from and d >= feed_from else "statistics/orders")
                                    + ", продажи — statistics/sales или отчёт о реализации")
    first_size_order = o["date"].min() if len(o) else end
    # до начала данных по размерам — итог дня по артикулу из воронки (без разбивки по размерам)
    f = funnel[(funnel["dt"] >= start) & (funnel["dt"] < first_size_order)]
    if len(f):
        add = pd.DataFrame({"Дата": f["dt"], "Размер": "все размеры",
                            "Заказано, шт": pd.to_numeric(f["ordersCount"]),
                            "Отменено, шт": pd.to_numeric(f["cancelCount"]),
                            "Источник": "итог по артикулу из воронки продаж (размеров WB не отдаёт); "
                                        "продано/возвраты — в строках по размерам"})
        df = pd.concat([df, add], ignore_index=True)
    cols = ["Дата", "Размер", "Заказано, шт", "Продано, шт", "Возврат, шт", "Отменено, шт", "Источник"]
    for c in cols[2:6]:
        if c not in df:
            df[c] = None
    df = df[cols]
    df.loc[df["Размер"] != "все размеры", ["Продано, шт", "Возврат, шт"]] = \
        df.loc[df["Размер"] != "все размеры", ["Продано, шт", "Возврат, шт"]].fillna(0)
    sized = (df["Размер"] != "все размеры") & (df["Дата"] >= first_size_order)
    df.loc[sized, ["Заказано, шт", "Отменено, шт"]] = df.loc[sized, ["Заказано, шт", "Отменено, шт"]].fillna(0)
    return _sort(df, ["Дата", "Размер"]), first_size_order


def cur_stock_sheet(nm, sidx):
    rows = {}
    for kind, url in (("WB", "wb-warehouses"), ("FBS", "seller-warehouses")):
        r = request("POST", f"{ANALYTICS}/api/analytics/v1/stocks-report/{url}", headers=wb._headers(),
                    body={"nmIds": [nm], "limit": 1000}, pause=21) or {}
        for it in (r.get("data") or {}).get("items") or []:
            size = sidx.get(int(it["chrtId"]), (0, "", str(it["chrtId"])))[2]
            x = rows.setdefault(size, {"На складах WB, шт": 0, "На своём складе (FBS), шт": 0,
                                       "В пути к клиенту, шт": 0, "В пути от клиента, шт": 0})
            if kind == "WB":
                x["На складах WB, шт"] += it.get("quantity", 0)
                x["В пути к клиенту, шт"] += it.get("inWayToClient", 0)
                x["В пути от клиента, шт"] += it.get("inWayFromClient", 0)
            else:
                x["На своём складе (FBS), шт"] += it.get("quantity", 0)
    df = pd.DataFrame([{"Дата снятия": dt.date.today().isoformat(), "Размер": k, **v} for k, v in rows.items()])
    return _sort(df, ["Размер"])


def stock_history(nm):
    end = yesterday()
    start = end - dt.timedelta(days=89)
    h = nm_report("STOCK_HISTORY_DAILY_CSV", {"nmIds": [nm], "currentPeriod": {"start": start.isoformat(), "end": end.isoformat()},
                                               "stockType": "", "skipDeletedNm": False})
    days = [c for c in h.columns if re.match(r"\d\d\.\d\d\.\d{4}$", c)]
    long = h.melt(id_vars=["SizeName", "OfficeName"], value_vars=days, var_name="d", value_name="q")
    long["Дата"] = pd.to_datetime(long["d"], format="%d.%m.%Y").dt.strftime("%Y-%m-%d")
    long["q"] = pd.to_numeric(long["q"], errors="coerce").fillna(0)
    return long.groupby(["SizeName", "Дата"])["q"].sum().reset_index().rename(columns={"SizeName": "Размер"})


def _after_launch(hist):
    """Оставляет дни начиная с первого дня, когда у размера был остаток: нули до появления размера
    (новые размеры 58, 60) — не дни без остатка."""
    first = hist[hist["q"] > 0].groupby("Размер")["Дата"].min()
    return hist[hist["Дата"] >= hist["Размер"].map(first).fillna("9999")]


def zero_days_sheet(hist):
    h = _after_launch(hist)
    z = h[h["q"] == 0][["Размер", "Дата"]].rename(columns={"Дата": "Дата, когда остаток был 0"})
    z["Остаток считается"] = "на 23:59, склады WB + свой склад"
    return _sort(z, ["Размер", "Дата, когда остаток был 0"])


def cur_summary_sheet(art, hist, o):
    s = read_csv(DATA / "wb_sales.csv")
    s = s[(s["article"].str.upper() == art) & (s["status"] == "sale")].assign(qty=lambda x: pd.to_numeric(x["qty"]))
    end = yesterday()
    rows = []
    for size in sorted(set(o["size"]) | set(s["size"]) | set(hist["Размер"]), key=_size_key):
        r = {"Размер": size}
        for n in (7, 14, 28):
            a = (end - dt.timedelta(days=n - 1)).isoformat()
            r[f"Заказов/день за {n} дн"] = round(o[(o["size"] == size) & (o["date"] >= a) & (o["date"] <= end.isoformat())]["qty"].sum() / n, 2)
        for n in (7, 14, 28):
            a = (end - dt.timedelta(days=n - 1)).isoformat()
            r[f"Выкупов/день за {n} дн"] = round(s[(s["size"] == size) & (s["date"] >= a) & (s["date"] <= end.isoformat())]["qty"].sum() / n, 2)
        a28 = (end - dt.timedelta(days=27)).isoformat()
        hz = _after_launch(hist)
        hz = hz[(hz["Размер"] == size) & (hz["Дата"] >= a28)]
        r["Дней без остатка за 28 дней"] = int((hz["q"] == 0).sum()) if len(hz) else None
        rows.append(r)
    return pd.DataFrame(rows)


# ---------- прошлый год ----------

def past_sales_sheet(funnel, real):
    r = real[real["vendor_code"].str.upper().isin(PAST) & real["doc_type"].isin(["Продажа", "Возврат"])].copy()
    r["Дата"] = r["sale_dt"].str[:10]
    r = r[(r["Дата"] >= PAST_FROM.isoformat()) & (r["Дата"] <= PAST_TO.isoformat())]
    r["q"] = pd.to_numeric(r["quantity"])
    r["Артикул"] = r["vendor_code"].str.upper()
    g = r.pivot_table(index=["Дата", "Артикул", "size"], columns="doc_type", values="q", aggfunc="sum", fill_value=0)
    g = g.rename(columns={"Продажа": "Продано, шт", "Возврат": "Возврат, шт"}).reset_index().rename(columns={"size": "Размер"})
    g["Заказано, шт"] = None
    g["Остаток на конец дня, шт"] = None
    g["Источник"] = "по размерам: отчёт о реализации (дата продажи)"
    f = funnel[funnel["Артикул"].isin(PAST) & (funnel["dt"] >= PAST_FROM.isoformat()) & (funnel["dt"] <= PAST_TO.isoformat())]
    add = pd.DataFrame({"Дата": f["dt"], "Артикул": f["Артикул"], "Размер": "все размеры",
                        "Заказано, шт": pd.to_numeric(f["ordersCount"]),
                        "Выкуплено из заказов дня, шт": pd.to_numeric(f["buyoutsCount"]),
                        "Отменено из заказов дня, шт": pd.to_numeric(f["cancelCount"]),
                        "Переходы в карточку": pd.to_numeric(f["openCardCount"]),
                        "Добавления в корзину": pd.to_numeric(f["addToCartCount"]),
                        "Источник": "итог по артикулу из воронки продаж (размеров WB не отдаёт)"})
    df = pd.concat([g, add], ignore_index=True)
    cols = ["Дата", "Артикул", "Размер", "Заказано, шт", "Продано, шт", "Остаток на конец дня, шт", "Возврат, шт",
            "Выкуплено из заказов дня, шт", "Отменено из заказов дня, шт", "Переходы в карточку",
            "Добавления в корзину", "Источник"]
    return _sort(df.reindex(columns=cols), ["Дата", "Артикул", "Размер"])


def past_prices_sheet(funnel, real):
    r = real[real["vendor_code"].str.upper().isin(PAST) & (real["doc_type"] == "Продажа")].copy()
    r["Дата"] = r["sale_dt"].str[:10]
    r = r[(r["Дата"] >= PAST_FROM.isoformat()) & (r["Дата"] <= PAST_TO.isoformat())]
    r["q"] = pd.to_numeric(r["quantity"])
    r = r[r["q"] > 0]
    r["buyer"] = pd.to_numeric(r["retail_amount"]) / r["q"]
    r["seller"] = pd.to_numeric(r["retail_price_with_disc"])
    r["Артикул"] = r["vendor_code"].str.upper()
    g = r.groupby(["Дата", "Артикул"]).agg(**{
        "Цена покупателя, руб (медиана по продажам дня)": ("buyer", "median"),
        "Цена продавца до СПП, руб (медиана по продажам дня)": ("seller", "median"),
        "Продаж в расчёте": ("q", "sum")}).round(0).reset_index()
    f = funnel[funnel["Артикул"].isin(PAST) & (funnel["dt"] >= PAST_FROM.isoformat()) & (funnel["dt"] <= PAST_TO.isoformat())].copy()
    f["n"] = pd.to_numeric(f["ordersCount"])
    f["Средняя цена заказа (воронка), руб"] = (pd.to_numeric(f["ordersSumRub"]) / f["n"].where(f["n"] > 0)).round(0)
    f = f.rename(columns={"dt": "Дата"})[["Дата", "Артикул", "Средняя цена заказа (воронка), руб"]]
    df = g.merge(f, on=["Дата", "Артикул"], how="outer")
    return df.sort_values(["Дата", "Артикул"])


def check(art, o, funnel):
    """Сверка с ЛК: заказы — воронка продаж (DETAIL_HISTORY_REPORT = «Аналитика → Воронка» в ЛК);
    продажи — детализация отчёта о реализации (= «Финансовые отчёты» в ЛК), там где недели закрыты."""
    end = yesterday()
    a = end - dt.timedelta(days=27)
    s = read_csv(DATA / "wb_sales.csv")
    o = o[(o["date"] >= a.isoformat()) & (o["date"] <= end.isoformat())]
    s = s[(s["article"].str.upper() == art) & (s["date"] >= a.isoformat()) & (s["date"] <= end.isoformat())]
    f = funnel[(funnel["Артикул"] == art) & (funnel["dt"] >= a.isoformat()) & (funnel["dt"] <= end.isoformat())]
    real = read_csv(RAW / "wb_realization.csv")
    real = real[(real["vendor_code"].str.upper() == art) & (real["doc_type"] == "Продажа")]
    last_week = real["report_to"].max()
    rs = real[(real["sale_dt"].str[:10] >= a.isoformat()) & (real["sale_dt"].str[:10] <= last_week)]
    ss = s[(s["status"] == "sale") & (s["date"] <= last_week)]
    rows = [
        ("Заказано, шт", f"{a}–{end}", "воронка продаж (ЛК)", pd.to_numeric(f["ordersCount"]).sum(), pd.to_numeric(o["qty"]).sum()),
        ("Продано, шт", f"{a}–{last_week}", "отчёт о реализации (ЛК)", pd.to_numeric(rs["quantity"]).sum(), pd.to_numeric(ss["qty"]).sum()),
    ]
    return pd.DataFrame([{"Показатель": n, "Период": p, "Эталон": src, "Эталон, шт": int(ref), "Выгрузка, шт": int(ours),
                          "Расхождение, шт": int(ours - ref), "Расхождение, %": round((ours - ref) / ref * 100, 2) if ref else None}
                         for n, p, src, ref, ours in rows])


def build():
    prices = json.loads((RAW / "wb_prices.json").read_text(encoding="utf-8"))
    nms = _nm_ids(prices, [CUR] + PAST)
    sidx = wb.size_index(prices)
    start_f = max(dt.date.today() - dt.timedelta(days=365), dt.date(2025, 10, 21))
    funnel = nm_report("DETAIL_HISTORY_REPORT", {"nmIDs": list(nms.values()), "startDate": start_f.isoformat(),
                                                 "endDate": yesterday().isoformat(), "timezone": "Europe/Moscow",
                                                 "aggregationLevel": "day", "skipDeletedNm": False})
    art_by_nm = {str(v): k for k, v in nms.items()}
    funnel["Артикул"] = funnel["nmID"].map(art_by_nm)
    real = pd.concat([x for x in (read_csv(RAW / "wb_realization_old_BZH.csv"), read_csv(RAW / "wb_realization.csv")) if x is not None])
    real = real.drop_duplicates("rrd_id")
    cur_real = real[(real["vendor_code"].str.upper() == CUR) & (real["doc_type"] == "Продажа")]
    first_sale = cur_real["sale_dt"].min()[:10]
    fo = funnel[(funnel["Артикул"] == CUR) & (pd.to_numeric(funnel["ordersCount"]) > 0)]["dt"].min()
    start = min(x for x in (first_sale, fo) if x)
    feed, feed_from = order_feed(nms[CUR], sidx)
    o = orders_by_size(CUR, feed, feed_from)
    orders, first_size = cur_orders_sheet(CUR, funnel[funnel["Артикул"] == CUR], start, o, feed_from)
    hist = stock_history(nms[CUR])
    sheets = {
        f"{CUR} — заказы и выкупы": orders,
        f"{CUR} — остатки": cur_stock_sheet(nms[CUR], sidx),
        f"{CUR} — сводка": cur_summary_sheet(CUR, hist, o),
        f"{CUR} — дни без остатка": zero_days_sheet(hist),
        "Прошлый год — продажи": past_sales_sheet(funnel, real),
        "Прошлый год — цены": past_prices_sheet(funnel, real),
    }
    chk = check(CUR, o, funnel)
    out = DATA / f"wb_plan_{CUR}_{dt.date.today().isoformat()}.xlsx"
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name[:31], index=False)
            ws = xw.sheets[name[:31]]
            for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 300)), 1):
                ws.column_dimensions[get_column_letter(i)].width = min(max(len(str(c.value or "")) for c in col) + 2, 45)
            ws.freeze_panes = "A2"
    chk.to_csv(out.with_name(out.stem + "_check.csv"), index=False, encoding="utf-8-sig")
    info = {"feed_from": feed_from, "start": start, "first_size_orders": first_size, "hist_from": hist["Дата"].min(), "hist_to": hist["Дата"].max()}
    log.info("Отчёт WB для плана: %s %s", out.name, info)
    return out, sheets, chk, info
