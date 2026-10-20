"""Сборка data/summary.xlsx: Продажи по артикулам и размерам, Остатки, Финансы."""
import datetime as dt

import pandas as pd
from openpyxl.utils import get_column_letter

from .common import log, period_of, read_csv, yesterday

SOLD = {"sale", "delivered"}
# группируем по модели и размеру: у WB один и тот же артикул встречается в разном регистре
KEYS = ["platform", "model", "size"]


def _articles(df):
    return df.groupby(KEYS)["article"].agg(lambda x: ", ".join(sorted(set(x))))


def _num(df, cols):
    for c in cols:
        if c in df:
            df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0)
    return df


def sales_sheet(data_dir):
    frames = []
    for name in ("wb_sales", "ozon_sales", "wb_orders", "ozon_orders"):
        df = read_csv(data_dir / f"{name}.csv")
        if df is not None and len(df):
            df["kind"] = "orders" if name.endswith("orders") else "sales"
            frames.append(df[KEYS + ["article", "date", "qty", "revenue", "status", "kind"]])
    if not frames:
        return pd.DataFrame(columns=KEYS)
    df = _num(pd.concat(frames, ignore_index=True), ["qty", "revenue"])
    d30 = (yesterday() - dt.timedelta(days=29)).isoformat()
    d7 = (yesterday() - dt.timedelta(days=6)).isoformat()
    s = df[df["kind"] == "sales"]
    o = df[df["kind"] == "orders"]
    sold, ret = s[s["status"].isin(SOLD)], s[s["status"] == "return"]
    canc = o[o["status"] == "cancelled"]
    out = pd.DataFrame(index=pd.MultiIndex.from_frame(df[KEYS].drop_duplicates()))
    out["articles"] = _articles(df)
    out["sold_total"] = sold.groupby(KEYS)["qty"].sum()
    out["returns_total"] = ret.groupby(KEYS)["qty"].sum()
    out["revenue_pre_spp_total"] = sold.groupby(KEYS)["revenue"].sum() - ret.groupby(KEYS)["revenue"].sum().reindex(out.index, fill_value=0)
    out["sold_30d"] = sold[sold["date"] >= d30].groupby(KEYS)["qty"].sum()
    out["sold_7d"] = sold[sold["date"] >= d7].groupby(KEYS)["qty"].sum()
    out["orders_30d"] = o[o["date"] >= d30].groupby(KEYS)["qty"].sum()
    out["cancelled_30d"] = canc[canc["date"] >= d30].groupby(KEYS)["qty"].sum()
    out["orders_7d"] = o[o["date"] >= d7].groupby(KEYS)["qty"].sum()
    # продано штук по периодам 16–15
    sold = sold.assign(period=sold["date"].map(period_of))
    per = sold.pivot_table(index=KEYS, columns="period", values="qty", aggfunc="sum")
    per.columns = [f"sold_{c}" for c in per.columns]
    out = out.join(per)
    num = [c for c in out.columns if c != "articles"]
    out[num] = out[num].fillna(0)
    return out.reset_index().sort_values(["model", "size", "platform"])


def stock_sheet(data_dir):
    def latest(prefix):
        files = sorted(data_dir.glob(f"{prefix}_stock_*.csv"))
        return read_csv(files[-1]) if files else None
    parts = []
    wb = latest("wb")
    if wb is not None and len(wb):
        wb = _num(wb, ["stock", "in_way_to_client", "in_way_from_client"])
        g = wb.groupby(KEYS + ["warehouse_type"])[["stock", "in_way_to_client", "in_way_from_client"]].sum()
        g = g.unstack("warehouse_type", fill_value=0)
        g.columns = [f"{a}_{b}" for a, b in g.columns]
        g = g.loc[:, (g != 0).any()]
        g.insert(0, "articles", _articles(wb))
        parts.append(g.reset_index().assign(date=wb["date"].iloc[0]))
    oz = latest("ozon")
    if oz is not None and len(oz):
        oz = _num(oz, ["free_to_sell", "reserved", "in_transit", "present"])
        tot = oz[oz["scheme"].str.endswith("_TOTAL")]
        g = tot.groupby(KEYS + ["scheme"])[["free_to_sell", "reserved"]].sum().unstack("scheme", fill_value=0)
        g.columns = [f"{a}_{b.replace('_TOTAL', '')}" for a, b in g.columns]
        tr = oz[oz["scheme"] == "FBO"].groupby(KEYS)["in_transit"].sum().rename("in_transit_FBO")
        g = g.join(tr).fillna(0)
        g.insert(0, "articles", _articles(tot))
        parts.append(g.reset_index().assign(date=oz["date"].iloc[0]))
    if not parts:
        return pd.DataFrame(columns=KEYS)
    out = pd.concat(parts, ignore_index=True)
    num = [c for c in out.columns if c not in KEYS + ["articles", "date"]]
    out[num] = out[num].fillna(0)
    out.insert(0, "date", out.pop("date"))
    return out.sort_values(["model", "size", "platform"])


def _autosize(ws):
    for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)), 1):
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 40)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def build(data_dir, finance_totals):
    path = data_dir / "summary.xlsx"
    sheets = {
        "Продажи по артикулам и размерам": sales_sheet(data_dir),
        "Остатки": stock_sheet(data_dir),
        "Финансы": finance_totals if finance_totals is not None else pd.DataFrame(),
    }
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name[:31], index=False)
            _autosize(xw.sheets[name[:31]])
    log.info("Записан data/summary.xlsx")
