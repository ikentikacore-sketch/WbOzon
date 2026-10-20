"""Разовый срез текущих цен и комиссий WB и Озон по выбранным артикулам (для плана отгрузок).

    python run.py --prices                       # БЖ2, БЖ4, БЖ5, БЖ11
    python run.py --prices "БЖ2,БЖ11,ШКЮ-1"

Результат: data/prices_commissions_<артикулы>_<дата>.xlsx (лист «Цены и комиссии» по ТЗ + лист «Детали») и .csv.

WB: цена после СПП = цена со скидкой продавца × (1 − СПП). СПП в API цен нет, берём медиану
поля spp по сегодняшним заказам артикула (если сегодня заказов нет — за 2 последних дня).
Комиссия — тариф по предмету (/api/v1/tariffs/commission), FBW = paidStorageKgvp.
Озон: цена продажи = цена продавца (marketing_seller_price, иначе price) из /v5/product/info/prices,
комиссия — commissions.sales_percent_fbo/fbs оттуда же.
"""
import datetime as dt
import re

import pandas as pd
from openpyxl.utils import get_column_letter

from . import ozon, wb
from .common import DATA, log, read_csv, request

DEFAULT = ["БЖ2", "БЖ4", "БЖ5", "БЖ11"]


def _size_key(size):
    m = re.match(r"\d+", str(size))
    return int(m.group()) if m else 0


def _wb(arts, today):
    prices = wb.load_prices()
    comm = request("GET", "https://common-api.wildberries.ru/api/v1/tariffs/commission",
                   headers=wb._headers(), params={"locale": "ru"})
    comm = {c["subjectName"]: c for c in (comm or {}).get("report", [])}
    since = (dt.date.today() - dt.timedelta(days=2)).isoformat()
    o = pd.DataFrame(request("GET", f"{wb.STAT}/api/v1/supplier/orders", headers=wb._headers(),
                             params={"dateFrom": since, "flag": 0}) or [])
    if len(o):
        o = o[o["supplierArticle"].str.upper().isin(arts)]
    st_files = sorted(DATA.glob("wb_stock_*.csv"))
    st = read_csv(st_files[-1]) if st_files else None
    real = read_csv(DATA / "raw" / "wb_realization.csv")
    kvw = None
    if real is not None:
        r = real[(real["doc_type"] == "Продажа") & real["vendor_code"].str.upper().isin(arts)]
        r = r[r["rr_date"] >= (dt.date.today() - dt.timedelta(days=14)).isoformat()]
        if len(r):
            kvw = float(pd.to_numeric(r["commission_percent"]).median())
    main, det = [], []
    for g in prices:
        art = g["vendorCode"].upper()
        if art not in arts:
            continue
        oo = o[o["nmId"] == g["nmID"]] if len(o) else o
        use = oo[oo["date"].str[:10] == today] if len(oo) else oo
        if not len(use):
            use = oo
        spp = float(use["spp"].median()) if len(use) else None
        subj = oo["subject"].iloc[0] if len(oo) else "Блузки"
        c = comm.get(subj, {})
        for s in g["sizes"]:
            size, pre = s["techSizeName"], float(s["discountedPrice"])
            after = round(pre * (1 - spp / 100)) if spp is not None else None
            sz = use[use["techSize"] == size] if len(use) else use
            stock = 0
            if st is not None:
                x = st[(st["article"].str.upper() == art) & (st["size"] == size)]
                stock = int(pd.to_numeric(x["stock"]).sum())
            main.append({"Площадка": "WB", "Артикул": art, "Размер": size, "Цена, руб": after,
                         "Комиссия, %": c.get("paidStorageKgvp"), "Дата снятия": today})
            det.append({
                "Площадка": "WB", "Артикул": art, "Размер": size, "nmId / SKU": g["nmID"],
                "Цена без скидки, руб": s["price"], "Скидка продавца, %": g["discount"],
                "Цена до СПП (база комиссии), руб": pre,
                "СПП (медиана по заказам), %": spp, "Заказов в расчёте СПП": len(use),
                "Цена после СПП, руб": after,
                "Факт. цена покупателя в заказах размера (медиана), руб":
                    float(sz["finishedPrice"].median()) if len(sz) else None,
                "Комиссия FBW (склад WB), %": c.get("paidStorageKgvp"),
                "Комиссия FBS, %": c.get("kgvpMarketplace"),
                "Комиссия в отчётах реализации за 14 дней (факт), %": kvw,
                "Остаток на складах, шт": stock, "Предмет": subj, "Дата снятия": today})
    return main, det


def _ozon(arts, today):
    op = ozon.load_prices()
    prod = ozon.load_products()
    st_files = sorted(DATA.glob("ozon_stock_*.csv"))
    ost = read_csv(st_files[-1]) if st_files else None
    buyer = {}
    for f in sorted(DATA.glob("ozon_finance_*.csv"))[-2:]:
        df = read_csv(f)
        df = df[(df["type_id"] == "sale") & (pd.to_numeric(df["amount"]) > 0)
                & (df["date"] >= (dt.date.today() - dt.timedelta(days=14)).isoformat())]
        for k, v in df.groupby("offer_id")["sale_price"]:
            buyer[k] = float(pd.to_numeric(v).median())
    main, det = [], []
    items = [(ozon.split_ozon_offer(o), o, v) for o, v in op.items()]
    for (base, size), offer, v in sorted(items, key=lambda x: (x[0][0], _size_key(x[0][1]))):
        if base.upper() not in arts:
            continue
        pr, cm = v["price"], v.get("commissions", {})
        price = float(pr.get("marketing_seller_price") or pr["price"])
        old = float(pr.get("old_price") or 0)
        free = 0
        if ost is not None:
            t = ost[(ost["article"] == offer) & ost["scheme"].str.endswith("_TOTAL")]
            free = int(pd.to_numeric(t["free_to_sell"]).sum())
        main.append({"Площадка": "Ozon", "Артикул": base.upper(), "Размер": size, "Цена, руб": price,
                     "Комиссия, %": cm.get("sales_percent_fbo"), "Дата снятия": today})
        det.append({
            "Площадка": "Ozon", "Артикул": base.upper(), "Размер": size, "nmId / SKU": ozon.sku_of(prod.get(offer, {})),
            "Цена без скидки, руб": old, "Скидка продавца, %": round((1 - float(pr["price"]) / old) * 100) if old else None,
            "Цена до СПП (база комиссии), руб": price,
            "Факт. цена покупателя за 14 дней (медиана, из начислений), руб": buyer.get(offer),
            "Комиссия FBO, %": cm.get("sales_percent_fbo"), "Комиссия FBS, %": cm.get("sales_percent_fbs"),
            "Эквайринг, руб/шт": v.get("acquiring"), "Остаток на складах, шт": free, "Дата снятия": today})
    return main, det


def build(arts=None):
    arts = [a.strip().upper() for a in (arts or DEFAULT) if a.strip()]
    today = dt.date.today().isoformat()
    main, det = [], []
    for name, fn in (("WB", _wb), ("Ozon", _ozon)):
        try:
            m, d = fn(arts, today)
            main += m
            det += d
        except Exception as e:  # noqa: BLE001 — одна площадка не должна останавливать другую
            log.error("Срез цен %s: %s", name, e)
    order = {a: i for i, a in enumerate(arts)}

    def tidy(df):
        if not len(df):
            return df
        key = df["Площадка"].map({"WB": 0, "Ozon": 1}) * 10000 + df["Артикул"].map(order).fillna(99) * 1000 \
            + df["Размер"].map(_size_key)
        return df.loc[key.sort_values(kind="stable").index].reset_index(drop=True)
    m, d = tidy(pd.DataFrame(main)), tidy(pd.DataFrame(det))
    tag = "_".join(arts) if len(arts) <= 6 else f"{len(arts)}art"
    out = DATA / f"prices_commissions_{tag}_{today}.xlsx"
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        m.to_excel(xw, sheet_name="Цены и комиссии", index=False)
        d.to_excel(xw, sheet_name="Детали", index=False)
        for ws in xw.sheets.values():
            for i, col in enumerate(ws.iter_cols(min_row=1, max_row=ws.max_row), 1):
                ws.column_dimensions[get_column_letter(i)].width = min(max(len(str(c.value or "")) for c in col) + 2, 40)
            ws.freeze_panes = "A2"
    m.to_csv(out.with_suffix(".csv"), index=False, encoding="utf-8-sig")
    log.info("Срез цен и комиссий: %s (%d строк)", out.name, len(m))
    return out, m, d
